"""Domain models and catalog for Taiwan active faults (TEM seismogenic structures)."""

from dataclasses import dataclass, field
from math import radians, sin, cos, sqrt, atan2
from pathlib import Path
from typing import Dict, List, Optional, Tuple
import openpyxl

from src.config import RAW_DATA_DIR


@dataclass(frozen=True)
class FaultParameter:
    """Attributes of an active seismogenic structure in Taiwan."""
    id: int
    name: str
    fault_type: str  # 'N' (Normal), 'R' (Reverse), 'SS' (Strike-Slip)
    rake: float
    dip: float
    depth_max_km: float
    mw_max: float
    slip_rate_mm_yr: float


@dataclass(frozen=True)
class FaultAlignment:
    """Geographic trace points for a fault structure."""
    id: int
    coordinates: List[Tuple[float, float]]  # List of (longitude, latitude)


@dataclass
class FaultDistanceResult:
    """Distance query result from a point to a fault trace."""
    fault: FaultParameter
    min_distance_km: float
    nearest_coord: Tuple[float, float]  # (lon, lat)


def haversine_distance_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    """Compute great-circle distance in kilometers using the Haversine formula."""
    r_earth_km = 6371.0
    phi1, phi2 = radians(lat1), radians(lat2)
    delta_phi = radians(lat2 - lat1)
    delta_lambda = radians(lon2 - lon1)

    a = sin(delta_phi / 2.0) ** 2 + cos(phi1) * cos(phi2) * sin(delta_lambda / 2.0) ** 2
    c = 2.0 * atan2(sqrt(a), sqrt(1.0 - a))
    return r_earth_km * c


class FaultCatalog:
    """In-memory catalog of all 38 on-land seismogenic structures in Taiwan."""

    def __init__(self, parameters: Dict[int, FaultParameter], alignments: Dict[int, FaultAlignment]):
        self._parameters = parameters
        self._alignments = alignments
        self._by_name: Dict[str, FaultParameter] = {
            f.name.lower().strip(): f for f in parameters.values()
        }

    @property
    def total_faults(self) -> int:
        return len(self._parameters)

    def get_by_id(self, fault_id: int) -> Optional[FaultParameter]:
        return self._parameters.get(fault_id)

    def get_by_name(self, name: str) -> Optional[FaultParameter]:
        return self._by_name.get(name.lower().strip())

    def get_alignment(self, fault_id: int) -> Optional[FaultAlignment]:
        return self._alignments.get(fault_id)

    def list_all(self) -> List[FaultParameter]:
        return sorted(self._parameters.values(), key=lambda f: f.id)

    def find_nearest_fault(self, target_lat: float, target_lon: float) -> Optional[FaultDistanceResult]:
        """Find the nearest fault trace to the target coordinates."""
        nearest_result: Optional[FaultDistanceResult] = None
        min_dist = float("inf")

        for fault_id, param in self._parameters.items():
            align = self._alignments.get(fault_id)
            if not align or not align.coordinates:
                continue

            for lon, lat in align.coordinates:
                dist = haversine_distance_km(target_lat, target_lon, lat, lon)
                if dist < min_dist:
                    min_dist = dist
                    nearest_result = FaultDistanceResult(
                        fault=param,
                        min_distance_km=round(dist, 2),
                        nearest_coord=(lon, lat)
                    )

        return nearest_result

    @classmethod
    def load_from_excel(
        cls,
        parameters_path: Optional[Path] = None,
        alignments_path: Optional[Path] = None
    ) -> "FaultCatalog":
        """Load and parse parameters and coordinate alignments from raw Excel assets."""
        p_path = parameters_path or (RAW_DATA_DIR / "Fault Parameters_update.xlsx")
        a_path = alignments_path or (RAW_DATA_DIR / "Fault Alignments.xlsx")

        # 1. Parse Fault Parameters
        params: Dict[int, FaultParameter] = {}
        wb_params = openpyxl.load_workbook(p_path, read_only=True)
        ws_params = wb_params.active
        for row in list(ws_params.iter_rows(values_only=True))[1:]:
            if row[0] is None:
                continue
            fault_id = int(row[0])
            params[fault_id] = FaultParameter(
                id=fault_id,
                name=str(row[1]).strip(),
                fault_type=str(row[2]).strip(),
                rake=float(row[3]),
                dip=float(row[4]),
                depth_max_km=float(row[5]),
                mw_max=float(row[6]),
                slip_rate_mm_yr=float(row[7]),
            )
        wb_params.close()

        # 2. Parse Fault Alignments
        alignments: Dict[int, FaultAlignment] = {}
        wb_align = openpyxl.load_workbook(a_path, read_only=True)
        ws_align = wb_align.active
        align_rows = list(ws_align.iter_rows(values_only=True))
        wb_align.close()

        header_row = align_rows[0]
        # Identify column pairs for each fault ID
        id_to_col_pair: Dict[int, Tuple[int, int]] = {}
        for col_idx in range(0, len(header_row), 2):
            val = header_row[col_idx]
            if val is not None:
                try:
                    f_id = int(val)
                    id_to_col_pair[f_id] = (col_idx, col_idx + 1)
                except (ValueError, TypeError):
                    pass

        # Extract coordinate points
        coord_map: Dict[int, List[Tuple[float, float]]] = {f_id: [] for f_id in id_to_col_pair}
        for row in align_rows[1:]:
            for f_id, (lon_col, lat_col) in id_to_col_pair.items():
                if lon_col < len(row) and lat_col < len(row):
                    lon_val = row[lon_col]
                    lat_val = row[lat_col]
                    if lon_val is not None and lat_val is not None:
                        try:
                            coord_map[f_id].append((float(lon_val), float(lat_val)))
                        except (ValueError, TypeError):
                            pass

        for f_id, coords in coord_map.items():
            alignments[f_id] = FaultAlignment(id=f_id, coordinates=coords)

        return cls(parameters=params, alignments=alignments)
